import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(12 if level == 1 else (8 if level == 2 else 6))
    h.paragraph_format.space_after = Pt(4)
    run = h.add_run(text)
    run.font.name = 'Times New Roman'
    if level == 1:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)
    elif level == 2:
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 60, 90)
    else:
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(50, 50, 50)
    return h

def render_table(doc, title, headers, data, col_widths, font_size=7.5, footnote=""):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(3)
    p_title.paragraph_format.keep_with_next = True
    r_title = p_title.add_run(title)
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(8.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 51, 102)

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_row = table.rows[0]
    hdr_row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for idx, header_text in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.width = Inches(col_widths[idx])
        set_cell_background(cell, "003366")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(header_text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size + 0.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, row_data in enumerate(data):
        row = table.rows[row_idx + 1]
        is_proposed = (len(row_data) > 0 and ("HDBO-MPA" in str(row_data[0]) or "Proposed" in str(row_data[0])))
        bg_color = "EBF2FA" if is_proposed else ("F7F9FC" if row_idx % 2 == 1 else "FFFFFF")
        for col_idx, cell_value in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.width = Inches(col_widths[col_idx])
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            
            val_str = str(cell_value)
            if len(val_str) <= 6 or val_str in ["✓", "✗", "100.0%"]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            run = p.add_run(val_str)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(font_size)
            if is_proposed or "✓" in val_str:
                run.font.bold = True
            if val_str == "✓":
                run.font.color.rgb = RGBColor(0, 128, 0)
            elif val_str == "✗":
                run.font.color.rgb = RGBColor(180, 0, 0)

    if footnote:
        p_fn = doc.add_paragraph()
        p_fn.paragraph_format.space_before = Pt(2)
        p_fn.paragraph_format.space_after = Pt(4)
        r_fn = p_fn.add_run(footnote)
        r_fn.font.name = 'Times New Roman'
        r_fn.font.size = Pt(7.5)
        r_fn.font.italic = True
        r_fn.font.color.rgb = RGBColor(80, 80, 80)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(3)

def add_figure_styled(doc, img_path, caption, width_inches=6.85):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(8)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(8.5)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(60, 60, 60)

def add_callout_box(doc, text_lines, title=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.85)
    set_cell_background(cell, "F2F4F7")
    set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if title:
        r_t = p.add_run(f"{title}\n")
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(9.5)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0, 51, 102)

    for i, line in enumerate(text_lines):
        if i > 0 or title:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
        r = p.add_run(line)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(8.5)
        if "ABSTRACT" in line or "Keywords" in line or line.startswith("Algorithm") or line.startswith("Input:") or line.startswith("Output:"):
            r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def add_body_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bp = p.add_run(bold_prefix)
        r_bp.font.name = 'Times New Roman'
        r_bp.font.size = Pt(9.5)
        r_bp.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9.5)
    return p
